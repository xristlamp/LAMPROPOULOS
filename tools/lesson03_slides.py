# Φτιάχνει τις διαφάνειες του Μαθήματος 3 (Λίστες και πλειάδες), σταδιακή έκδοση.
# Χρήση:  py tools\lesson03_slides.py
# Παίρνει την πρώτη διαφάνεια και το θέμα (SAEK) από το παλιό αρχείο και γράφει:
#   Desktop\ΣΑΕΚ_GLYFADA\Μάθημα03_Λίστες_και_πλειάδες_v2.pptx
# Μετά: tools\convert.ps1 για PDF. Αν αλλάξει η σειρά των διαφανειών, διόρθωσε το REFS στο pylab.js.
import math, os, re, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn

SITE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(SITE, "files", "glyfada", "python-03-listes.pptx")
OUT = os.path.join(os.path.expanduser("~"), "Desktop", "ΣΑΕΚ_GLYFADA", "Μάθημα03_Λίστες_και_πλειάδες_v2.pptx")
if len(sys.argv) > 1: OUT = sys.argv[1]

# ---------- χρώματα & διαστάσεις (ίδια με τα άλλα μαθήματα) ----------
INK, TEXT, MUTED, LINE, PAPER = "171B26", "4A5265", "8A91A3", "D3D9E4", "FBFCFE"
BLUE, BLUE_BG, GOLD, GOLD_BG = "2F4BD8", "E4E8FB", "9A6B0C", "F6EDD6"
SANS, MONO = "Segoe UI", "Consolas"
L, R = 566928, 11494008
W = R - L
COL_L, COL_W = L, 5394960            # αριστερή στήλη
COL_R, COL_RW = 6236208, R - 6236208 # δεξιά στήλη
TOP = 1417320
FOOT_Y = 6053328

prs = Presentation(SRC)
LAYOUT = [l for l in prs.slide_layouts if l.name == "SAEK"][0]

# κράτα μόνο την 1η διαφάνεια (τίτλος)
ids = prs.slides._sldIdLst
for sid in list(ids)[1:]:
    prs.part.drop_rel(sid.rId)
    ids.remove(sid)


# ---------- βοηθοί ----------
def rgb(c): return RGBColor.from_string(c)

def run(p, text, size, color=INK, font=SANS, bold=False, spc=None):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.name = Pt(size), bold, font
    f.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        rPr.append(rPr.makeelement(qn(tag), {"typeface": font}))
    if spc: rPr.set("spc", str(spc))
    return r

def rich(p, text, size, color=INK, font=SANS, bold_color=None):
    """**έντονο** μέσα στο κείμενο"""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part: run(p, part, size, (bold_color or color) if i % 2 else color, font, bold=bool(i % 2))

def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    t = s.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf

def shape(s, x, y, w, h, fill=None, line=None, kind=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(kind, x, y, w, h)
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = 9525
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.05
    return sh

def est_h(text, w, size, mono=False, lh=1.22):
    """εκτίμηση ύψους κειμένου που αναδιπλώνεται"""
    plain = text.replace("**", "")
    per = max(8, (w / 12700) / (size * (0.6 if mono else 0.53)))
    lines = sum(max(1, math.ceil(len(seg) / per)) for seg in plain.split("\n"))
    return int(lines * size * lh * 12700)

def label(s, x, y, w, text, color=BLUE):
    tf = box(s, x, y, w, 201168, MSO_ANCHOR.MIDDLE)
    run(tf.paragraphs[0], text, 9, color, MONO, bold=True, spc=100)

def kicker_title(s, num, kicker, title, title_w=9692640):
    tf = box(s, L, 384048, 8778240, 237744, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, f"{num:02d}  ", 10, BLUE, MONO, spc=140)
    run(p, kicker, 10, MUTED, MONO, spc=140)
    tf = box(s, 530352, 658368, title_w, 566928)
    run(tf.paragraphs[0], title, 30, INK, SANS, bold=True)

def lab_badge(s, nums):
    """ετικέτα πάνω δεξιά: σε ποια άσκηση του εργαστηρίου πάει η διαφάνεια"""
    w, h, x, y = 1737360, 502920, R - 1737360, 694944
    shape(s, x, y, w, h, GOLD_BG, None, MSO_SHAPE.ROUNDED_RECTANGLE)
    shape(s, x, y, 36576, h, GOLD)
    tf = box(s, x + 137160, y + 54864, w - 182880, h - 109728, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, "ΣΤΟ ΕΡΓΑΣΤΗΡΙΟ", 8, GOLD, MONO, bold=True, spc=80)
    p = tf.add_paragraph()
    txt = ("Άσκηση " if len(nums) == 1 else "Ασκήσεις ") + " · ".join(str(n) for n in nums)
    run(p, txt, 13, INK, SANS, bold=True)

def subtitle(s, text, y=1261872):
    tf = box(s, L, y, W, 274320, MSO_ANCHOR.MIDDLE)
    rich(tf.paragraphs[0], text, 12, TEXT, SANS)

def footer(s, tag, text):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, L, FOOT_Y, L + 11119104, FOOT_Y)
    ln.line.color.rgb = rgb(LINE); ln.line.width = 9525; ln.line.dash_style = MSO_LINE.DASH
    tf = box(s, L, 6089904, 10241280, 329184, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, tag + "   ", 9, GOLD, MONO, bold=True, spc=100)
    rich(p, text, 10.5, MUTED, SANS)

KW = set("for in if elif else while break True False None and or not import".split())
BI = set("print input int float str len range sum max min sorted round list tuple abs".split())
TOK = re.compile(r'(#.*$)|(f?"[^"]*")|([A-Za-z_]\w*)|(\s+)|(.)')

def code_line(p, line, size):
    if line.startswith("»"):                              # γραμμή εξόδου
        out, _, com = line[1:].partition("   #")
        run(p, out, size, TEXT, MONO)
        if com: run(p, "   #" + com, size, MUTED, MONO)
        return
    for m in TOK.finditer(line):
        com, st, word, sp, ch = m.groups()
        if com: run(p, com, size, MUTED, MONO)
        elif st: run(p, st, size, GOLD, MONO)
        elif word: run(p, word, size, BLUE if (word in KW or word in BI) else INK, MONO)
        else: run(p, sp or ch, size, INK, MONO)
    if not line: run(p, " ", size, INK, MONO)

def code(s, x, y, w, src, bar=BLUE, size=11, tag=None, tag_color=BLUE, h=None):
    """κουτί κώδικα· οι γραμμές που ξεκινούν με » είναι αυτό που εμφανίζεται. Επιστρέφει το κάτω όριο."""
    if tag:
        label(s, x, y, w, tag, tag_color)
        y += 283464
    lines = src.strip("\n").split("\n")
    lh = int(size * 1.2 * 12700)
    hh = h or len(lines) * lh + 2 * 118872
    shape(s, x, y, w, hh, PAPER, LINE)
    shape(s, x, y, 36576, hh, bar)
    tf = box(s, x + 164592, y + 109728, w - 246888, hh - 182880)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(size * 1.2)
        code_line(p, line, size)
    return y + hh

def card(s, x, y, w, title, body, tone="blue", h=None, size=11):
    fill = BLUE_BG if tone == "blue" else GOLD_BG
    bh = est_h(body, w - 329184, size)
    hh = h or (109728 + 274320 + 54864 + bh + 128016)
    shape(s, x, y, w, hh, fill, None, MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box(s, x + 164592, y + 109728, w - 329184, 274320, MSO_ANCHOR.MIDDLE)
    run(tf.paragraphs[0], title, 13, INK, SANS, bold=True)
    tf = box(s, x + 164592, y + 438912, w - 329184, bh)
    for i, part in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        rich(p, part, size, TEXT, SANS, bold_color=INK)
    return y + hh

def cards(s, x, y, w, items, gap=137160):
    for it in items:
        y = card(s, x, y, w, *it) + gap
    return y

def prompt(s, x, y, w, text, tag="ΕΚΦΩΝΗΣΗ", size=11.5):
    bh = est_h(text, w - 219456, size)
    hh = 73152 + 201168 + 36576 + bh + 109728
    shape(s, x, y, w, hh, BLUE_BG, BLUE_BG)
    shape(s, x, y, 41148, hh, BLUE)
    label(s, x + 164592, y + 73152, w - 219456, tag)
    tf = box(s, x + 164592, y + 292608, w - 219456, bh)
    rich(tf.paragraphs[0], text, size, INK, SANS)
    return y + hh

def _borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = tcPr.makeelement(qn(side), {"w": "9525" if side == "a:lnB" else "0", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        if side == "a:lnB":
            sf = ln.makeelement(qn("a:solidFill"), {}); c = sf.makeelement(qn("a:srgbClr"), {"val": LINE}); sf.append(c); ln.append(sf)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.insert(i, ln)

def table(s, x, y, colw, rows, mono=(), size=10.5, row_h=338328, head=True, hl=None):
    """rows[0] = επικεφαλίδες. mono: στήλες σε Consolas. hl: {(γραμμή, στήλη): χρώμα}"""
    gs = s.shapes.add_table(len(rows), len(colw), x, y, sum(colw), row_h * len(rows))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")): tblPr.remove(el)
    for i, cw in enumerate(colw): tbl.columns[i].width = cw
    for r in range(len(rows)):
        tbl.rows[r].height = row_h if (r or not head) else 320040
        for c in range(len(colw)):
            cell = tbl.cell(r, c)
            cell.fill.background()
            _borders(cell)
            cell.margin_left = cell.margin_right = 76200
            cell.margin_top = cell.margin_bottom = 45720
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            txt = str(rows[r][c])
            if r == 0 and head:
                run(p, txt.upper(), 9, MUTED, MONO, bold=True, spc=60)
            else:
                col = (hl or {}).get((r, c), INK)
                rich(p, txt, size, col, MONO if c in mono else SANS)
    return y + sum(tbl.rows[r].height for r in range(len(rows)))

def new(num, kicker, title, lab=None, title_w=None):
    s = prs.slides.add_slide(LAYOUT)
    kicker_title(s, num, kicker, title, title_w or (8686800 if lab else 9692640))
    if lab: lab_badge(s, lab)
    return s


# ======================= ΔΙΑΦΑΝΕΙΕΣ =======================
STEPS = 15
def step(n, name): return f"ΒΗΜΑ {n} ΑΠΟ {STEPS} · {name}"

# 1 · τίτλος (από το παλιό αρχείο, αλλάζει μόνο η γραμμή με τα στοιχεία)
first = prs.slides[0]
for sh in first.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("Θεωρία"):
        r0 = sh.text_frame.paragraphs[0].runs[0]
        r0.text = "Θεωρία σε 15 μικρά βήματα + 18 ασκήσεις στο εργαστήριο          Διάρκεια: 4 ώρες (3η–6η)          27 διαφάνειες"

# 2 · στόχοι
s = new(2, "ΣΤΟΧΟΙ", "Τι θα ξέρεις στο τέλος")
goals = [
    "Να φτιάχνεις μια **λίστα** και να παίρνεις ένα στοιχείο από τη **θέση** του.",
    "Να **αλλάζεις**, να **προσθέτεις** και να **αφαιρείς** στοιχεία.",
    "Να περνάς από όλα τα στοιχεία με **for** και να βρίσκεις άθροισμα, μέσο όρο, μέγιστο.",
    "Να **γεμίζεις** μια λίστα με τιμές που δίνει ο χρήστης και να την **ταξινομείς**.",
    "Να χρησιμοποιείς **πλειάδα** για τιμές που πάνε μαζί, π.χ. μια θέση (x, y) — θα τη χρειαστείς στο pygame.",
]
y = 1481328
for i, g in enumerate(goals, 1):
    shape(s, L, y, 457200, 237744, PAPER, LINE)
    tf = box(s, L, y, 457200, 237744, MSO_ANCHOR.MIDDLE)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run(tf.paragraphs[0], str(i), 8.5, BLUE, MONO)
    tf = box(s, 1152144, y - 18288, 10341864, 274320, MSO_ANCHOR.MIDDLE)
    rich(tf.paragraphs[0], g, 13, INK, SANS)
    y += 384048
y += 182880
card(s, L, y, 5303520, "Πώς δουλεύουμε",
     "Η θεωρία είναι χωρισμένη σε **15 μικρά βήματα**. Κάθε βήμα έχει πάνω δεξιά την ετικέτα **ΣΤΟ ΕΡΓΑΣΤΗΡΙΟ** "
     "με τον αριθμό της άσκησης στο site.\n**Διαβάζεις τη διαφάνεια → λύνεις την άσκηση → πας στο επόμενο βήμα.**", h=1600200)
card(s, COL_R - 91440, y, COL_RW + 91440, "Από το εύκολο στο δύσκολο",
     "Ασκήσεις **1–7**: εύκολες, μία εντολή τη φορά.\n"
     "Ασκήσεις **8–14**: μέτριες, συνδυάζουν 2–3 βήματα.\n"
     "Ασκήσεις **15–17**: δύσκολες, όλα μαζί. Η **18** είναι για το σπίτι.", tone="gold", h=1600200)
footer(s, "ΣΥΝΔΕΣΗ", "Το for της προηγούμενης φοράς περνούσε από αριθμούς (range) και από γράμματα. Σήμερα θα περνά από **λίστες**.")

# 3 · ζέσταμα
s = new(3, "ΖΕΣΤΑΜΑ", "Θυμάσαι από την προηγούμενη φορά;")
y = prompt(s, L, TOP, W, "Πέντε γρήγορες ερωτήσεις, 5 λεπτά. Τι εμφανίζει κάθε κομμάτι; Πρώτα σκέψου, μετά δες τη λύση.")
label(s, L, y + 182880, 3000000, "ΛΥΣΗ", GOLD)
table(s, L, y + 457200, [3474720, 1463040, 2011680], [
    ["Κώδικας", "Αποτέλεσμα", "Γιατί"],
    ['for i in range(2, 6): print(i)', "2 3 4 5", "χωρίς το τέλος"],
    ['for ch in "ok": print(ch)', "o k", "γράμμα-γράμμα"],
    ['len("Python")', "6", "πόσα γράμματα"],
    ['"Python"[0]', '"P"', "το πρώτο γράμμα"],
    ['"y" in "Python"', "True", "υπάρχει μέσα"],
], mono=(0, 1))
card(s, 7790688, y + 457200, 3703320, "Γιατί αυτά",
     "Ό,τι ξέρεις για τα κείμενα — **len**, **[0]**, **in**, **for** — δουλεύει **σχεδόν ίδιο** και στις λίστες. Ξεκινάμε με πλεονέκτημα.", h=1554480)
footer(s, "ΠΑΡΑΤΗΡΗΣΗ", "Ένα κείμενο είναι μια σειρά από **γράμματα**. Μια λίστα είναι μια σειρά από **ό,τι θέλεις**: αριθμούς, λέξεις…")

# 4 · το πρόβλημα
s = new(4, "ΤΟ ΠΡΟΒΛΗΜΑ", "30 βαθμοί = 30 μεταβλητές;")
b1 = code(s, COL_L, TOP, COL_W, '''
grade1 = 12
grade2 = 8.5
grade3 = 17
# ... άλλες 27 γραμμές ...
avg = (grade1 + grade2 + grade3 + ...) / 30
''', bar=GOLD, tag="ΧΩΡΙΣ ΛΙΣΤΑ", tag_color=GOLD)
b2 = code(s, COL_R, TOP, COL_RW, '''
grades = [12, 8.5, 17, 9, 15]

avg = sum(grades) / len(grades)
print(avg)
»12.3
''', tag="ΜΕ ΛΙΣΤΑ")
y = max(b1, b2) + 228600
card(s, COL_L, y, COL_W, "Τι πάει στραβά",
     "Όσο περισσότεροι οι βαθμοί, τόσο μεγαλύτερος ο κώδικας. Κι αν τους δίνει ο χρήστης, **δεν ξέρεις** πόσες μεταβλητές να φτιάξεις.", tone="gold", h=1371600)
card(s, COL_R, y, COL_RW, "Τι κερδίζεις",
     "**Ένα όνομα** για όλες τις τιμές. Ο ίδιος κώδικας δουλεύει για 5, 30 ή 3.000 βαθμούς, χωρίς καμία αλλαγή.", h=1371600)
footer(s, "Η ΙΔΕΑ", "Η λίστα είναι σαν ένα **κουτί με θήκες**: ένα όνομα, πολλές τιμές, η καθεμία στη θέση της.")

# 5 · βήμα 1
s = new(5, step(1, "ΛΙΣΤΑ"), "Φτιάχνω μια λίστα", lab=[1])
code(s, COL_L, TOP, COL_W, '''
grades = [12, 8.5, 17]
names = ["Άννα", "Νίκος", "Ελένη"]
empty = []

print(names)
»['Άννα', 'Νίκος', 'Ελένη']
print(len(names))
»3
print(len(empty))
»0
''')
cards(s, COL_R, TOP, COL_RW, [
    ("Πώς γράφεται", "Ανοίγεις αγκύλη **[**, γράφεις τις τιμές **χωρισμένες με κόμμα**, κλείνεις με **]**. Τα κείμενα θέλουν εισαγωγικά, οι αριθμοί όχι."),
    ("Πόσα στοιχεία έχει;", "Το **len(λίστα)** δίνει το πλήθος — όπως το len(\"Python\") έδινε πόσα γράμματα έχει η λέξη."),
    ("Άδεια λίστα", "Το **[]** είναι μια λίστα χωρίς τίποτα μέσα. Θα τη γεμίσουμε αργότερα (βήμα 10)."),
])
footer(s, "ΘΥΜΗΣΟΥ", "Αγκύλες **[ ]** = λίστα. Ένα στοιχείο ή εκατό, πάντα με **κόμμα** ανάμεσα.")

# 6 · βήμα 2
s = new(6, step(2, "ΘΕΣΕΙΣ"), "Παίρνω ένα στοιχείο από τη θέση του", lab=[2])
subtitle(s, 'fruits = ["μήλο", "αχλάδι", "κεράσι", "σύκο"]')
table(s, L, 1691640, [2148840, 2194560, 2194560, 2194560, 2194560], [
    ["Θέση", "0", "1", "2", "3"],
    ["τιμή", '"μήλο"', '"αχλάδι"', '"κεράσι"', '"σύκο"'],
    ["από το τέλος", "-4", "-3", "-2", "-1"],
], mono=(1, 2, 3, 4))
code(s, COL_L, 2834640, COL_W, '''
print(fruits[0])
»μήλο
print(fruits[2])
»κεράσι
print(fruits[-1])
»σύκο
''')
cards(s, COL_R, 2834640, COL_RW, [
    ("Το μέτρημα ξεκινά από το 0", "Το **πρώτο** στοιχείο είναι στη θέση **0**, το δεύτερο στη **1**… Με 4 στοιχεία η τελευταία θέση είναι η **3**."),
    ("Το [-1] είναι το τελευταίο", "Δεν χρειάζεται να ξέρεις πόσα στοιχεία έχει η λίστα: το **[-1]** είναι πάντα το τελευταίο, το **[-2]** το προτελευταίο."),
])
footer(s, "ΣΥΧΝΟ ΛΑΘΟΣ", "**fruits[4]** → IndexError: ζήτησες θέση που **δεν υπάρχει**. Εδώ οι θέσεις είναι 0, 1, 2, 3.")

# 7 · βήμα 3
s = new(7, step(3, "ΑΛΛΑΓΗ"), "Αλλάζω ένα στοιχείο", lab=[3])
b = code(s, COL_L, TOP, COL_W, '''
grades = [12, 8, 17]
grades[1] = 14      # ο 2ος βαθμός γίνεται 14
print(grades)
»[12, 14, 17]
''')
code(s, COL_L, b + 228600, COL_W, '''
grades[0] = grades[0] + 1   # +1 στον 1ο
print(grades)
»[13, 14, 17]
''', tag="ΜΕ ΠΡΑΞΗ")
cards(s, COL_R, TOP, COL_RW, [
    ("Όπως μια μεταβλητή", "Όπως γράφεις **x = 5**, γράφεις **grades[1] = 14**: «βάλε στη θέση 1 την τιμή 14». Η παλιά τιμή χάνεται."),
    ("Ποια θέση;", "Οι θέσεις μετράνε από το 0. Ο **2ος** βαθμός είναι στη θέση **1**."),
    ("Στο κείμενο δεν γίνεται", "Το word[0] = \"Μ\" σε κείμενο βγάζει σφάλμα. **Η λίστα αλλάζει, το κείμενο όχι.**", "gold"),
])
footer(s, "ΠΡΟΣΟΧΗ", "Αλλάζεις μόνο θέση που **υπάρχει**. Για νέο στοιχείο θέλεις **append** — επόμενο βήμα.")

# 8 · βήμα 4
s = new(8, step(4, "ΠΡΟΣΘΗΚΗ"), "Προσθέτω στοιχεία", lab=[4])
code(s, COL_L, TOP, COL_W, '''
cart = ["ψωμί"]

cart.append("γάλα")       # στο τέλος
print(cart)
»['ψωμί', 'γάλα']

cart.insert(0, "αυγά")    # στη θέση 0
print(cart)
»['αυγά', 'ψωμί', 'γάλα']
''')
cards(s, COL_R, TOP, COL_RW, [
    ("append(τιμή)", "Βάζει την τιμή **στο τέλος** της λίστας. Αυτό θα χρησιμοποιείς σχεδόν πάντα."),
    ("insert(θέση, τιμή)", "Βάζει την τιμή **στη θέση** που λες. Τα υπόλοιπα πάνε μία θέση δεξιά."),
    ("Η τελεία", "Γράφεις **όνομα_λίστας.append(…)**. Η λίστα αλλάζει αμέσως, δεν χρειάζεται «cart = …»."),
])
footer(s, "ΣΥΧΝΟ ΛΑΘΟΣ", "**cart = cart.append(\"τυρί\")** → χάνεις τη λίστα! Γράψε σκέτο cart.append(\"τυρί\").")

# 9 · βήμα 5
s = new(9, step(5, "ΑΦΑΙΡΕΣΗ"), "Αφαιρώ στοιχεία", lab=[5])
code(s, COL_L, TOP, COL_W, '''
cart = ["αυγά", "ψωμί", "γάλα", "τυρί"]

cart.remove("ψωμί")    # βγάζει την τιμή
print(cart)
»['αυγά', 'γάλα', 'τυρί']

last = cart.pop()      # βγάζει το τελευταίο
print(last)
»τυρί
print(cart)
»['αυγά', 'γάλα']
''')
cards(s, COL_R, TOP, COL_RW, [
    ("remove(τιμή)", "Ψάχνει την **τιμή** και τη βγάζει από τη λίστα (την πρώτη που θα βρει)."),
    ("pop()", "Βγάζει το **τελευταίο** στοιχείο και σου το δίνει, για να το κρατήσεις σε μεταβλητή. Το **pop(0)** βγάζει το πρώτο."),
])
footer(s, "ΣΥΧΝΟ ΛΑΘΟΣ", "**cart.remove(\"μέλι\")** όταν δεν υπάρχει μέλι → ValueError. Έλεγξε πρώτα με **in** — επόμενο βήμα.")

# 10 · βήμα 6
s = new(10, step(6, "ΑΝΑΖΗΤΗΣΗ"), "Υπάρχει στη λίστα;", lab=[6])
b = code(s, COL_L, TOP, COL_W, '''
guests = ["Άννα", "Νίκος", "Ελένη"]
print("Νίκος" in guests)
»True
print("Μαρία" in guests)
»False
''')
code(s, COL_L, b + 228600, COL_W, '''
name = input("Όνομα: ")
if name in guests:
    print("Καλώς ήρθες!")
else:
    print("Δεν είσαι στη λίστα")
''', tag="ΜΕ IF")
cards(s, COL_R, TOP, COL_RW, [
    ("in", "Απαντά **True** ή **False**: υπάρχει η τιμή μέσα στη λίστα; Είναι το ίδιο in που έψαχνε γράμμα σε κείμενο."),
    ("Αφαίρεση χωρίς σφάλμα", "**if \"μέλι\" in cart:** και από κάτω **cart.remove(\"μέλι\")**. Αν δεν υπάρχει, απλώς δεν γίνεται τίποτα."),
    ("Πόσες φορές;", "Το **grades.count(10)** μετράει πόσες φορές υπάρχει το 10 στη λίστα."),
])
footer(s, "ΣΥΝΔΕΣΗ", "Το if … in … είναι το ίδιο if του Μαθήματος 2. Η συνθήκη τώρα ρωτάει μια λίστα.")

# 11 · άσκηση στο χαρτί Α
s = new(11, "ΑΣΚΗΣΗ ΣΤΟ ΧΑΡΤΙ Α", "Τι έχει η λίστα μετά από κάθε γραμμή;")
y = prompt(s, L, TOP, W, "Γράψε στο τετράδιο πώς είναι η λίστα **μετά από κάθε γραμμή**. Μετά τρέξε τον κώδικα στο «🧪 Ελεύθερο» του εργαστηρίου και έλεγξε.")
label(s, L, y + 182880, 3000000, "ΛΥΣΗ", GOLD)
table(s, L, y + 457200, [2651760, 2560320, 1737360], [
    ["Γραμμή", "Η λίστα μετά", "Τι έγινε"],
    ["lst = [3, 1, 2]", "[3, 1, 2]", "φτιάχτηκε"],
    ["lst.append(5)", "[3, 1, 2, 5]", "5 στο τέλος"],
    ["lst.insert(0, 9)", "[9, 3, 1, 2, 5]", "9 στη θέση 0"],
    ["lst[1] = 7", "[9, 7, 1, 2, 5]", "η θέση 1 έγινε 7"],
    ["lst.remove(1)", "[9, 7, 2, 5]", "έφυγε η τιμή 1"],
    ["x = lst.pop()", "[9, 7, 2]", "x = 5"],
], mono=(0, 1))
card(s, 7790688, y + 457200, 3703320, "Η παγίδα",
     "Το **remove(1)** βγάζει την **τιμή 1** — όχι το στοιχείο που είναι στη θέση 1 (εκείνη τη στιγμή το 7).", tone="gold", h=1600200)
footer(s, "ΤΕΧΝΙΚΗ", "Μία γραμμή του πίνακα για κάθε εντολή — όπως ο πίνακας τιμών του Μαθήματος 2.")

# 12 · βήμα 7
s = new(12, step(7, "FOR"), "Περνάω από όλα τα στοιχεία με for", lab=[7])
b = code(s, COL_L, TOP, COL_W, '''
names = ["Άννα", "Νίκος", "Ελένη"]
for name in names:
    print("Γεια σου,", name)
»Γεια σου, Άννα
»Γεια σου, Νίκος
»Γεια σου, Ελένη
''')
label(s, COL_R, TOP, COL_RW, "ΓΥΡΟ-ΓΥΡΟ")
table(s, COL_R, TOP + 283464, [1188720, 1737360, 2331720], [
    ["Γύρος", "name", "Τυπώνει"],
    ["1ος", '"Άννα"', "Γεια σου, Άννα"],
    ["2ος", '"Νίκος"', "Γεια σου, Νίκος"],
    ["3ος", '"Ελένη"', "Γεια σου, Ελένη"],
], mono=(1,))
card(s, COL_R, 3200400, COL_RW, "Το όνομα το διαλέγεις εσύ",
     "**for name in names** · **for g in grades** · **for x in nums**. Συνήθως: η λίστα στον πληθυντικό, η μεταβλητή στον ενικό.")
card(s, COL_L, b + 228600, COL_W, "Πώς δουλεύει",
     "Σε κάθε γύρο η μεταβλητή **name** παίρνει το **επόμενο στοιχείο**. Όταν τελειώσουν τα στοιχεία, ο βρόχος σταματά μόνος του.")
footer(s, "ΙΔΙΟ ΜΕ ΠΡΙΝ", "Το for ch in \"ok\" περνούσε από γράμματα. Το **for name in names** περνά από στοιχεία.")

# 13 · βήμα 8
s = new(13, step(8, "ΕΤΟΙΜΕΣ ΑΠΑΝΤΗΣΕΙΣ"), "len · sum · max · min", lab=[8])
subtitle(s, "grades = [12, 8.5, 17, 9, 15]")
b = table(s, L, 1691640, [3200400, 1371600, 2377440], [
    ["Γράφεις", "Παίρνεις", "Τι είναι"],
    ["len(grades)", "5", "πόσα στοιχεία"],
    ["sum(grades)", "61.5", "άθροισμα"],
    ["max(grades)", "17", "ο μεγαλύτερος"],
    ["min(grades)", "8.5", "ο μικρότερος"],
    ["sum(grades) / len(grades)", "12.3", "**μέσος όρος**"],
], mono=(0, 1))
code(s, L, b + 228600, 6949440, '''
avg = sum(grades) / len(grades)
print(f"Μέσος όρος: {avg:.2f}")
»Μέσος όρος: 12.30
''')
cards(s, 7790688, 1691640, 3703320, [
    ("Μέσος όρος", "Δεν υπάρχει έτοιμη εντολή. Είναι πάντα **άθροισμα / πλήθος**, δηλαδή **sum / len**."),
    ("Και στις λέξεις", "max([\"Άννα\", \"Ελένη\"]) → **'Ελένη'**: η τελευταία αλφαβητικά.", "gold"),
])
footer(s, "ΘΥΜΗΣΟΥ", "Το **:.2f** (Μάθημα 1) δείχνει 2 δεκαδικά, το **:.1f** ένα.")

# 14 · βήμα 9
s = new(14, step(9, "ΜΕΤΡΑΩ ΚΑΙ ΔΙΑΛΕΓΩ"), "for + if: πόσοι; ποιοι;", lab=[9, 10])
b1 = code(s, COL_L, TOP, COL_W, '''
grades = [12, 8.5, 17, 9, 15]
count = 0
for g in grades:
    if g >= 10:
        count = count + 1
print("Πέρασαν:", count)
»Πέρασαν: 3
''', tag="1 · ΜΕΤΡΗΤΗΣ: ΠΟΣΟΙ ΠΕΡΑΣΑΝ;")
b2 = code(s, COL_R, TOP, COL_RW, '''
grades = [12, 8.5, 17, 9, 15]
passed = []
for g in grades:
    if g >= 10:
        passed.append(g)
print(passed)
»[12, 17, 15]
''', tag="2 · ΝΕΑ ΛΙΣΤΑ: ΠΟΙΟΙ ΠΕΡΑΣΑΝ;")
y = max(b1, b2) + 228600
card(s, COL_L, y, COL_W, "Μετρητής", "Ξεκινά από **0 πριν** τον βρόχο. Μέσα, **+1** κάθε φορά που ισχύει το if. Όπως ο αθροιστής του Μαθήματος 2.", h=1188720)
card(s, COL_R, y, COL_RW, "Νέα λίστα", "Ξεκινά **άδεια [] πριν** τον βρόχο. Μέσα, **append** ό,τι περνά το if.", h=1188720)
footer(s, "ΤΟ ΜΟΤΙΒΟ", "**Πριν** τον βρόχο: 0 ή []. **Μέσα**: if και πρόσθεση. **Μετά** τον βρόχο: print.")

# 15 · βήμα 10
s = new(15, step(10, "ΓΕΜΙΣΜΑ"), "Γεμίζω τη λίστα από τον χρήστη", lab=[11, 18])
b1 = code(s, COL_L, TOP, COL_W, '''
nums = []
for i in range(3):
    x = int(input("Αριθμός: "))
    nums.append(x)
print(nums)
»Αριθμός: 4
»Αριθμός: 7
»Αριθμός: 1
»[4, 7, 1]
''', tag="1 · ΞΕΡΩ ΠΟΣΕΣ ΤΙΜΕΣ: FOR")
b2 = code(s, COL_R, TOP, COL_RW, '''
words = []
while True:
    w = input("Λέξη: ")
    if w == "":
        break
    words.append(w)
print(words)
»Λέξη: γάτα
»Λέξη: σκύλος
»Λέξη:              # σκέτο Enter
»['γάτα', 'σκύλος']
''', tag="2 · ΔΕΝ ΞΕΡΩ ΠΟΣΕΣ: WHILE")
y = max(b1, b2) + 228600
card(s, COL_L, y, COL_W, "for", "Όταν ξέρεις **πόσες** τιμές θα δοθούν (ή τις ρωτάς πρώτα: «Πόσοι μαθητές;»).", h=1005840)
card(s, COL_R, y, COL_RW, "while True + break", "Όταν σταματά ο χρήστης: πατά σκέτο **Enter** (κενή γραμμή) και το break βγαίνει από τον βρόχο.", h=1005840)
footer(s, "ΣΥΧΝΟ ΛΑΘΟΣ", "Το **nums = []** μπαίνει **πριν** τον βρόχο. Αν το βάλεις μέσα, η λίστα αδειάζει σε κάθε γύρο!")

# 16 · βήμα 11
s = new(16, step(11, "ΤΑΞΙΝΟΜΗΣΗ"), "Βάζω τη λίστα σε σειρά", lab=[12])
b1 = code(s, COL_L, TOP, COL_W, '''
nums = [5, 2, 9, 1]
nums.sort()
print(nums)
»[1, 2, 5, 9]
nums.sort(reverse=True)
print(nums)
»[9, 5, 2, 1]
''', tag="NUMS.SORT() · ΑΛΛΑΖΕΙ ΤΗ ΛΙΣΤΑ")
b2 = code(s, COL_R, TOP, COL_RW, '''
nums = [5, 2, 9, 1]
new = sorted(nums)
print(new)
»[1, 2, 5, 9]
print(nums)
»[5, 2, 9, 1]   # η παλιά έμεινε ίδια
''', tag="SORTED(NUMS) · ΦΤΙΑΧΝΕΙ ΝΕΑ")
y = max(b1, b2) + 228600
card(s, COL_L, y, COL_W, "reverse=True", "Από τον **μεγαλύτερο** στον μικρότερο. Χωρίς αυτό: από τον μικρότερο στον μεγαλύτερο.", h=1005840)
card(s, COL_R, y, COL_RW, "Λέξεις", "Ταξινομούνται **αλφαβητικά**: sorted([\"Νίκος\", \"Άννα\"]) → ['Άννα', 'Νίκος'].", h=1005840)
footer(s, "ΣΥΧΝΟ ΛΑΘΟΣ", "**nums = nums.sort()** → η λίστα γίνεται None (τίποτα). Το sort() γράφεται μόνο του.")

# 17 · βήμα 12
s = new(17, step(12, "ΚΟΜΜΑΤΙ"), "Παίρνω ένα κομμάτι της λίστας", lab=[12])
subtitle(s, "nums = [10, 20, 30, 40, 50, 60]")
b = table(s, L, 1691640, [1828800, 1920240, 3200400], [
    ["Γράφεις", "Παίρνεις", "Με λόγια"],
    ["nums[0:3]", "[10, 20, 30]", "από τη θέση 0 ως **πριν** την 3"],
    ["nums[:3]", "[10, 20, 30]", "τα 3 πρώτα"],
    ["nums[2:5]", "[30, 40, 50]", "θέσεις 2, 3, 4"],
    ["nums[3:]", "[40, 50, 60]", "από τη θέση 3 ως το τέλος"],
    ["nums[-2:]", "[50, 60]", "τα 2 τελευταία"],
], mono=(0, 1))
code(s, L, b + 228600, 6949440, '''
grades = [12, 17, 9, 19, 15]
top = sorted(grades, reverse=True)
print(top[:3])
»[19, 17, 15]
''', tag="ΟΙ 3 ΚΑΛΥΤΕΡΟΙ")
cards(s, 7790688, 1691640, 3703320, [
    ("Ο κανόνας", "**[αρχή:τέλος]**: η αρχή μετράει, το **τέλος όχι**. Ίδιο με το range(αρχή, τέλος)."),
    ("Τα 3 καλύτερα", "Ταξινόμησε από τον μεγαλύτερο (βήμα 11) και πάρε το **[:3]**.", "gold"),
])
footer(s, "ΚΑΙ ΣΤΑ ΚΕΙΜΕΝΑ", "Το **\"Python\"[:2]** δίνει \"Py\".")

# 18 · βήμα 13
s = new(18, step(13, "ΔΥΟ ΛΙΣΤΕΣ"), "Θέση και τιμή μαζί", lab=[16])
b = code(s, COL_L, TOP, COL_W, '''
names  = ["Άννα", "Νίκος", "Ελένη"]
grades = [17, 9, 15]

for i in range(len(names)):
    print(names[i], grades[i])
»Άννα 17
»Νίκος 9
»Ελένη 15
''')
label(s, COL_R, TOP, COL_RW, "ΓΥΡΟ-ΓΥΡΟ")
table(s, COL_R, TOP + 283464, [1005840, 2011680, 2240280], [
    ["i", "names[i]", "grades[i]"],
    ["0", '"Άννα"', "17"],
    ["1", '"Νίκος"', "9"],
    ["2", '"Ελένη"', "15"],
], mono=(0, 1, 2))
card(s, COL_R, 3200400, COL_RW, "Ποιος έχει τον μεγαλύτερο;",
     "**grades.index(17)** → 0: σε ποια θέση είναι το 17.\nΜετά **names[0]** → \"Άννα\".", tone="gold")
card(s, COL_L, b + 228600, COL_W, "range(len(names))",
     "Δίνει τις θέσεις **0, 1, 2…**, όσα και τα στοιχεία. Με το **ίδιο i** παίρνεις από τις δύο λίστες το όνομα και τον βαθμό του ίδιου μαθητή.")
footer(s, "ΠΡΟΣΟΧΗ", "Οι δύο λίστες πρέπει να έχουν **την ίδια σειρά** και **το ίδιο μήκος**.")

# 19 · συχνά λάθη
s = new(19, "ΠΡΟΣΟΧΗ", "Τα πιο συχνά λάθη με τις λίστες")
table(s, L, TOP, [3383280, 2286000, 5257800], [
    ["Τι έγραψες", "Τι βγαίνει", "Η διόρθωση"],
    ["fruits[4]   (με 4 στοιχεία)", "IndexError", "οι θέσεις είναι 0–3· για το τελευταίο: **fruits[-1]**"],
    ['cart.remove("μέλι")', "ValueError", 'έλεγξε πρώτα: **if "μέλι" in cart:**'],
    ["nums = nums.sort()", "None", "σκέτο **nums.sort()**"],
    ["nums = []   (μέσα στο for)", "αδειάζει κάθε γύρο", "βάλε το **nums = []** πριν το for"],
    ['["μήλο" "αχλάδι"]', "['μήλοαχλάδι']", "ξέχασες το **κόμμα**"],
    ["sum(x) / len(x)   (άδεια x)", "ZeroDivisionError", "έλεγξε πρώτα: **if len(x) > 0:**"],
], mono=(0, 1), row_h=420624, hl={(r, 1): "B42318" for r in range(1, 7)})
card(s, L, 4663440, W, "Διάβασε το μήνυμα του σφάλματος",
     "Η **τελευταία γραμμή** λέει **τι** πήγε στραβά (π.χ. IndexError). Το **line N** λέει **σε ποια γραμμή**. Στο εργαστήριο, ο έλεγχος σου το εξηγεί και στα ελληνικά.", h=1097280)
footer(s, "ΣΥΜΒΟΥΛΗ", "Όταν κάτι δεν βγαίνει, βάλε **print(λίστα)** μετά από κάθε βήμα και δες πώς αλλάζει.")

# 20 · βήμα 14
s = new(20, step(14, "ΠΛΕΙΑΔΑ"), "Πλειάδα: τιμές που πάνε μαζί", lab=[13])
b = code(s, COL_L, TOP, COL_W, '''
point = (3, 4)        # θέση x, y
red = (255, 0, 0)     # χρώμα R, G, B

print(point[0])
»3
x, y = point          # μοιράζω σε 2 μεταβλητές
print(x, y)
»3 4
''')
code(s, COL_L, b + 228600, COL_W, '''
point[0] = 10
»TypeError: 'tuple' object does not support item assignment
''', bar=GOLD, size=10, tag="ΔΕΝ ΑΛΛΑΖΕΙ", tag_color=GOLD)
cards(s, COL_R, TOP, COL_RW, [
    ("Παρενθέσεις αντί για αγκύλες", "Η πλειάδα (στα αγγλικά tuple) γράφεται με **( )**. Θέσεις, **len**, **in** και **for** δουλεύουν όπως στη λίστα."),
    ("Δεν αλλάζει", "Δεν έχει append ούτε remove, ούτε αλλάζει θέση. Τη φτιάχνεις μία φορά και μένει έτσι."),
    ("x, y = point", "Οι τιμές **μοιράζονται με τη σειρά** στις μεταβλητές αριστερά. Όσες τιμές, τόσες μεταβλητές."),
])
footer(s, "PYGAME", "Στο pygame οι θέσεις είναι **(x, y)** και τα χρώματα **(R, G, B)**: screen.fill((255, 255, 255)).")

# 21 · άσκηση στο χαρτί Β
s = new(21, "ΑΣΚΗΣΗ ΣΤΟ ΧΑΡΤΙ Β", "Λίστα ή πλειάδα;")
y = prompt(s, L, TOP, W, "Τι θα διάλεγες σε κάθε περίπτωση; Πρώτα σκέψου: θα **αλλάζει** ή μένει **ίδιο**;")
label(s, L, y + 182880, 3000000, "ΛΥΣΗ", GOLD)
table(s, L, y + 457200, [3566160, 1188720, 2286000], [
    ["Περίπτωση", "Διαλέγω", "Γιατί"],
    ["α) Τα προϊόντα στο καλάθι ενός e-shop", "**λίστα**", "ο πελάτης βάζει και βγάζει"],
    ["β) Η θέση (x, y) του παίκτη", "**πλειάδα**", "δύο τιμές που πάνε μαζί"],
    ["γ) Ένα χρώμα RGB", "**πλειάδα**", "έτσι το θέλει το pygame"],
    ["δ) Οι βαθμοί που δίνει ο χρήστης", "**λίστα**", "γεμίζει με append"],
    ["ε) Οι 7 μέρες της εβδομάδας", "**πλειάδα**", "δεν αλλάζουν ποτέ"],
])
card(s, 7790688, y + 457200, 3703320, "Ο γρήγορος κανόνας",
     "Θα βάζεις ή θα βγάζεις στοιχεία; → **λίστα**.\nΕίναι ένα «πακέτο» που μένει ίδιο; → **πλειάδα**.", h=1554480)
footer(s, "ΣΗΜΕΙΩΣΗ", "Και η λίστα θα «δούλευε» στα β, γ, ε. Η πλειάδα όμως **προστατεύει** τις τιμές από αλλαγή κατά λάθος.")

# 22 · βήμα 15
s = new(22, step(15, "ΠΙΝΑΚΑΣ"), "Λίστα με λίστες: πίνακας", lab=[14, 17])
b = code(s, COL_L, TOP, COL_W, '''
m = [[1, 2, 3],
     [4, 5, 6],
     [7, 8, 9]]

print(m[1])
»[4, 5, 6]
print(m[1][2])
»6
''')
code(s, COL_L, b + 228600, COL_W, '''
for row in m:
    print(sum(row))
»6
»15
»24
''', tag="ΑΘΡΟΙΣΜΑ ΚΑΘΕ ΓΡΑΜΜΗΣ")
label(s, COL_R, TOP, COL_RW, "ΤΟ m ΣΑΝ ΠΙΝΑΚΑΣ")
table(s, COL_R, TOP + 283464, [1554480, 1234440, 1234440, 1234440], [
    ["", "στήλη 0", "στήλη 1", "στήλη 2"],
    ["γραμμή 0", "1", "2", "3"],
    ["γραμμή 1", "4", "5", "**6**"],
    ["γραμμή 2", "7", "8", "9"],
], mono=(1, 2, 3), hl={(2, 3): BLUE})
cards(s, COL_R, 3108960, COL_RW, [
    ("Δύο θέσεις", "Πρώτα η **γραμμή**, μετά η **στήλη**: **m[γραμμή][στήλη]**. Το m[1][2] είναι η γραμμή 1, στήλη 2 → 6."),
    ("Πού θα το δεις", "Τρίλιζα, σκακιέρα, η πίστα ενός παιχνιδιού 2D: όλα είναι πίνακες.", "gold"),
])
footer(s, "ΠΛΗΘΗ", "**len(m)** = πόσες γραμμές · **len(m[0])** = πόσες στήλες.")

# 23 · οι ασκήσεις του εργαστηρίου
s = new(23, "ΕΡΓΑΣΤΗΡΙΟ", "Οι 18 ασκήσεις, βήμα-βήμα")
EX = [
    (1, "Η πρώτη μου λίστα", "Εύκολο", "5"), (2, "Πρώτη, τρίτη, τελευταία", "Εύκολο", "6"),
    (3, "Διόρθωση βαθμού", "Εύκολο", "7"), (4, "Καλάθι: βάζω προϊόντα", "Εύκολο", "8"),
    (5, "Καλάθι: βγάζω προϊόντα", "Εύκολο", "9"), (6, "Είσαι στη λίστα;", "Εύκολο", "10"),
    (7, "Γεια σε όλους", "Εύκολο", "12"), (8, "Θερμοκρασίες εβδομάδας", "Μέτριο", "13"),
    (9, "Πόσοι πέρασαν;", "Μέτριο", "14"), (10, "Μόνο οι ζυγοί", "Μέτριο", "14"),
    (11, "Γέμισε τη λίστα", "Μέτριο", "15"), (12, "Οι 3 καλύτεροι", "Μέτριο", "16, 17"),
    (13, "Πλειάδα: πόσο μακριά;", "Μέτριο", "20"), (14, "Πίνακας: πόντοι ομάδων", "Μέτριο", "22"),
    (15, "Βαθμολογίες τάξης", "Δύσκολο", "24"), (16, "Πάνω από τον μέσο όρο", "Δύσκολο", "18, 25"),
    (17, "Μαγικό τετράγωνο", "Ο στόχος", "22, 26"), (18, "Λέξεις μέχρι κενή γραμμή", "Για το σπίτι", "15, 16"),
]
LVL = {"Εύκολο": "1B7F4B", "Μέτριο": GOLD, "Δύσκολο": "B42318", "Ο στόχος": "B42318", "Για το σπίτι": BLUE}
for k, part in enumerate((EX[:9], EX[9:])):
    rows = [["#", "Άσκηση", "Επίπεδο", "Διαφάνεια"]] + [[str(n), t, lv, d] for n, t, lv, d in part]
    hl = {(i + 1, 2): LVL[part[i][2]] for i in range(len(part))}
    table(s, COL_L if k == 0 else COL_R, TOP, [457200, 2651760, 1234440, 1051560 if k == 0 else 914400], rows, mono=(0, 3), size=10.5, row_h=393192, hl=hl)
footer(s, "ΣΤΟ SITE", "Εργαστήριο → **Μάθημα 3**. Κάθε άσκηση έχει το «📖 Στις διαφάνειες» που ανοίγει κατευθείαν τη σωστή διαφάνεια.")

# 24 · δύσκολο 1
s = new(24, "ΕΡΓΑΣΤΗΡΙΟ · ΔΥΣΚΟΛΟ", "Βαθμολογίες τάξης", lab=[15])
y = prompt(s, L, TOP, W, "Διάβασε το πλήθος των μαθητών και τους βαθμούς τους. Εμφάνισε **μέσο όρο**, **μέγιστο** και **ελάχιστο**, **πόσοι πέρασαν** (≥ 10) και τους βαθμούς **ταξινομημένους** από τον μεγαλύτερο.")
label(s, L, y + 137160, 5029200, "ΛΥΣΗ · ΔΟΚΙΜΑΣΕ ΠΡΩΤΑ ΜΟΝΟΣ/Η ΣΟΥ", GOLD)
code(s, L, y + 411480, 6035040, '''
n = int(input("Πόσοι μαθητές; "))
grades = []
for i in range(n):
    g = float(input(f"Βαθμός {i + 1}: "))
    grades.append(g)

avg = sum(grades) / len(grades)
passed = 0
for g in grades:
    if g >= 10:
        passed = passed + 1

print(f"Μέσος όρος: {avg:.2f}")
print("Μέγιστος:", max(grades), " Ελάχιστος:", min(grades))
print("Πέρασαν:", passed, "από", n)
print("Ταξινομημένοι:", sorted(grades, reverse=True))
''', bar=GOLD, size=9.5)
b = code(s, 6876288, y + 411480, 4617720, '''
»Πόσοι μαθητές; 5
»Βαθμός 1: 12
»Βαθμός 2: 8.5
»...
»Μέσος όρος: 12.30
»Μέγιστος: 17.0  Ελάχιστος: 8.5
»Πέρασαν: 3 από 5
»Ταξινομημένοι: [17.0, 15.0, 12.0, 9.0, 8.5]
''', size=9.5)
card(s, 6876288, b + 137160, 4617720, "Από ποια βήματα;",
     "**10** γέμισμα · **8** sum, max, min · **9** μετρητής · **11** sorted")
footer(s, "ΓΙΑΤΙ 17.0", "Διαβάσαμε με **float()**, οπότε όλοι οι βαθμοί είναι δεκαδικοί και τυπώνονται με .0.")

# 25 · δύσκολο 2
s = new(25, "ΕΡΓΑΣΤΗΡΙΟ · ΔΥΣΚΟΛΟ", "Ποιοι είναι πάνω από τον μέσο όρο;", lab=[16])
y = prompt(s, L, TOP, W, "Έχεις δύο λίστες: τα **ονόματα** και τους **βαθμούς** (στην ίδια σειρά). Εμφάνισε τον μέσο όρο, όσους είναι **πάνω** από αυτόν και το όνομα του **καλύτερου**.")
label(s, L, y + 137160, 5029200, "ΛΥΣΗ · ΔΟΚΙΜΑΣΕ ΠΡΩΤΑ ΜΟΝΟΣ/Η ΣΟΥ", GOLD)
code(s, L, y + 411480, 6035040, '''
names = ["Άννα", "Νίκος", "Μαρία", "Γιώργος", "Ελένη"]
grades = [12, 8.5, 17, 9, 15]

avg = sum(grades) / len(grades)
print(f"Μέσος όρος: {avg:.2f}")

for i in range(len(names)):
    if grades[i] > avg:
        print(" ", names[i], grades[i])

best = grades.index(max(grades))
print("Καλύτερος/η:", names[best])
''', bar=GOLD, size=10)
b = code(s, 6876288, y + 411480, 4617720, '''
»Μέσος όρος: 12.30
»  Μαρία 17
»  Ελένη 15
»Καλύτερος/η: Μαρία
''', size=10)
cards(s, 6876288, b + 137160, 4617720, [
    ("Η γραμμή του best", "max(grades) → **17** · grades.index(17) → **2** · names[2] → **Μαρία**", "gold"),
    ("Από ποια βήματα;", "**8** sum, len, max · **13** δύο λίστες με range(len(…))"),
])
footer(s, "ΣΥΜΒΟΥΛΗ", "Κόλλησες; Λύσε πρώτα μόνο τον μέσο όρο, μετά μόνο τη λίστα, και στο τέλος τον καλύτερο.")

# 26 · ο στόχος
s = new(26, "ΕΡΓΑΣΤΗΡΙΟ · Ο ΣΤΟΧΟΣ", "Είναι μαγικό τετράγωνο;", lab=[17])
y = prompt(s, L, TOP, W, "Ένας πίνακας 3×3 είναι **μαγικός** όταν όλες οι γραμμές, όλες οι στήλες και οι δύο διαγώνιοι έχουν **το ίδιο άθροισμα**. Μάζεψε τα αθροίσματα σε μια λίστα και πες αν είναι μαγικός.")
label(s, L, y + 137160, 5029200, "ΛΥΣΗ · ΔΟΚΙΜΑΣΕ ΠΡΩΤΑ ΜΟΝΟΣ/Η ΣΟΥ", GOLD)
code(s, L, y + 411480, 6035040, '''
m = [[2, 7, 6],
     [9, 5, 1],
     [4, 3, 8]]

sums = []
for row in m:                              # γραμμές
    sums.append(sum(row))
for c in range(3):                         # στήλες
    sums.append(m[0][c] + m[1][c] + m[2][c])
sums.append(m[0][0] + m[1][1] + m[2][2])   # διαγώνιος ↘
sums.append(m[0][2] + m[1][1] + m[2][0])   # διαγώνιος ↙

print("Αθροίσματα:", sums)
if min(sums) == max(sums):
    print("Μαγικό!")
else:
    print("Όχι μαγικό")
''', bar=GOLD, size=9.5)
b = code(s, 6876288, y + 411480, 4617720, '''
»Αθροίσματα: [15, 15, 15, 15, 15, 15, 15, 15]
»Μαγικό!
''', size=9.5)
cards(s, 6876288, b + 137160, 4617720, [
    ("Η ιδέα", "Αν ο **μικρότερος** από τα αθροίσματα είναι ίσος με τον **μεγαλύτερο**, τότε είναι **όλα ίδια**."),
    ("Από ποια βήματα;", "**15** πίνακας · **4** append · **8** sum, min, max", "gold"),
])
footer(s, "ΠΡΟΚΛΗΣΗ", "Άλλαξε το 5 σε 6 και τρέξε ξανά. Τι βγαίνει; Γιατί;")

# 27 · σύνοψη
s = new(27, "ΣΥΝΟΨΗ", "Τι κρατάμε")
table(s, L, TOP, [3017520, 3931920], [
    ["Θέλω να…", "Γράφω"],
    ["φτιάξω λίστα", "nums = [4, 7, 1]"],
    ["πάρω το πρώτο / το τελευταίο", "nums[0] · nums[-1]"],
    ["αλλάξω ένα στοιχείο", "nums[1] = 9"],
    ["προσθέσω στο τέλος", "nums.append(5)"],
    ["βγάλω μια τιμή", "nums.remove(7) · nums.pop()"],
    ["δω αν υπάρχει", "if 7 in nums:"],
    ["περάσω από όλα", "for x in nums:"],
    ["βρω πλήθος, άθροισμα, max, min", "len · sum · max · min"],
    ["ταξινομήσω", "nums.sort() · sorted(nums)"],
    ["πάρω τα 3 πρώτα", "nums[:3]"],
], mono=(1,), row_h=374904)
cards(s, 7790688, TOP, 3703320, [
    ("Πλειάδα ( )", "Σαν λίστα που **δεν αλλάζει**. Για τιμές που πάνε μαζί: **(x, y)**, **(R, G, B)**. x, y = point."),
    ("Επόμενη φορά · Δευ 26/10", "**Σύνολα και λεξικά**: συλλογές χωρίς διπλότυπα και αναζήτηση με «κλειδί», π.χ. όνομα → τηλέφωνο.", "gold"),
])
footer(s, "ΓΙΑ ΤΟ ΣΠΙΤΙ", "Άσκηση **18** στο εργαστήριο: διάβαζε λέξεις μέχρι κενή γραμμή (βήμα 10).")

assert len(prs.slides) == 27, len(prs.slides)
prs.save(OUT)
print("OK", OUT, len(prs.slides), "διαφάνειες")
