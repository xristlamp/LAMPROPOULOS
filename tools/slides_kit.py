# Κοινά εργαλεία για τις διαφάνειες των μαθημάτων Python (ίδιο στυλ με τα υπόλοιπα μαθήματα).
# Χρήση σε ένα lessonNN_slides.py:
#   from slides_kit import *
#   start("python-04-syn-lex.pptx")      # κρατά την 1η διαφάνεια και το θέμα του παλιού αρχείου
#   s = new(2, "ΣΤΟΧΟΙ", "Τι θα ξέρεις στο τέλος") ...
#   finish("Μάθημα04_..._v2.pptx", 24)
import copy, math, os, re, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
SITE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------- χρώματα & διαστάσεις (ίδια με τα άλλα μαθήματα) ----------
INK, TEXT, MUTED, LINE, PAPER = "171B26", "4A5265", "8A91A3", "D3D9E4", "FBFCFE"
BLUE, BLUE_BG, GOLD, GOLD_BG = "2F4BD8", "E4E8FB", "9A6B0C", "F6EDD6"
SANS, MONO = "Segoe UI", "Consolas"
L, R = 566928, 11494008
W = R - L
COL_L, COL_W = L, 5394960            # αριστερή στήλη
COL_R, COL_RW = 6236208, R - 6236208 # δεξιά στήλη
TOP = 1417320
FOOT_Y = 6053328

prs = LAYOUT = None
CHROME = []


def start(src_name, folder="glyfada"):
    """Ανοίγει το παλιό pptx του μαθήματος (από files/<folder>, ή πλήρες path) και κρατά μόνο την 1η διαφάνεια.
    Αν το αρχείο δεν έχει layout «SAEK» (π.χ. Μαρκόπουλο), αντιγράφει το πλαίσιο (μπάρα, λογότυπο, σχολή, όνομα)
    από τη 2η διαφάνεια σε κάθε νέα διαφάνεια."""
    global prs, LAYOUT
    prs = Presentation(os.path.join(SITE, "files", folder, src_name))
    saek = [l for l in prs.slide_layouts if l.name == "SAEK"]
    LAYOUT = saek[0] if saek else prs.slides[0].slide_layout
    CHROME.clear()
    if not saek and len(prs.slides) > 1:
        src = prs.slides[1]
        bg = src._element.cSld.bg
        if bg is not None: CHROME.append((copy.deepcopy(bg), "bg"))
        for sh in list(src.shapes)[:4]:                   # μπάρα, λογότυπο, σχολή, όνομα
            img = src.part.related_part(sh._element.blipFill.blip.rEmbed) if sh.shape_type == 13 else None   # 13 = εικόνα
            CHROME.append((copy.deepcopy(sh._element), img))
    ids = prs.slides._sldIdLst
    for sid in list(ids)[1:]:
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    return prs


def title_meta(text):
    """Αλλάζει τη γραμμή «Θεωρία + … · Διάρκεια · N διαφάνειες» της 1ης διαφάνειας."""
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("Θεωρία"):
            sh.text_frame.paragraphs[0].runs[0].text = text


def title_slide(top, title, sub, meta, foot):
    """Ξαναγράφει όλη την 1η διαφάνεια (για νέο μάθημα που ξεκινά από το αρχείο ενός άλλου)."""
    texts = [sh for sh in prs.slides[0].shapes if sh.has_text_frame]
    for sh, text in zip(texts, (top, title, sub, meta, foot)):
        sh.text_frame.paragraphs[0].runs[0].text = text


def finish(out_name, count, folder="ΣΑΕΚ_GLYFADA"):
    """Γράφει στο Desktop/<folder> (ή στο path του 1ου ορίσματος)."""
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser("~"), "Desktop", folder, out_name)
    assert len(prs.slides) == count, len(prs.slides)
    prs.save(out)
    print("OK", out, len(prs.slides), "διαφάνειες")

# ---------- βοηθοί ----------
def rgb(c): return RGBColor.from_string(c)

def run(p, text, size, color=INK, font=SANS, bold=False, spc=None):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.name = Pt(size), bold, font
    f.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        rPr.append(rPr.makeelement(qn(tag), {"typeface": font}))
    if spc: rPr.set("spc", str(spc))
    return r

def rich(p, text, size, color=INK, font=SANS, bold_color=None):
    """**έντονο** μέσα στο κείμενο · \\*\\* = σκέτα αστεράκια (π.χ. \\*\\*kwargs)"""
    for i, part in enumerate(re.split(r"(?<!\\)\*\*", text)):
        part = part.replace("\\*", "*")
        if part: run(p, part, size, (bold_color or color) if i % 2 else color, font, bold=bool(i % 2))

def box(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    t = s.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf

def shape(s, x, y, w, h, fill=None, line=None, kind=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(kind, x, y, w, h)
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb = rgb(line); sh.line.width = 9525
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = 0.05
    return sh

def est_h(text, w, size, mono=False, lh=1.22):
    """εκτίμηση ύψους κειμένου που αναδιπλώνεται"""
    plain = text.replace("**", "")
    per = max(8, (w / 12700) / (size * (0.6 if mono else 0.53)))
    lines = sum(max(1, math.ceil(len(seg) / per)) for seg in plain.split("\n"))
    return int(lines * size * lh * 12700)

def label(s, x, y, w, text, color=BLUE):
    tf = box(s, x, y, w, 201168, MSO_ANCHOR.MIDDLE)
    run(tf.paragraphs[0], text, 9, color, MONO, bold=True, spc=100)

def kicker_title(s, num, kicker, title, title_w=9692640):
    tf = box(s, L, 384048, 8778240, 237744, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, f"{num:02d}  ", 10, BLUE, MONO, spc=140)
    run(p, kicker, 10, MUTED, MONO, spc=140)
    tf = box(s, 530352, 658368, title_w, 566928)
    run(tf.paragraphs[0], title, 30, INK, SANS, bold=True)

def lab_badge(s, nums):
    """ετικέτα πάνω δεξιά: σε ποια άσκηση του εργαστηρίου πάει η διαφάνεια"""
    w, h, x, y = 1737360, 502920, R - 1737360, 694944
    shape(s, x, y, w, h, GOLD_BG, None, MSO_SHAPE.ROUNDED_RECTANGLE)
    shape(s, x, y, 36576, h, GOLD)
    tf = box(s, x + 137160, y + 54864, w - 182880, h - 109728, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, "ΣΤΟ ΕΡΓΑΣΤΗΡΙΟ", 8, GOLD, MONO, bold=True, spc=80)
    p = tf.add_paragraph()
    txt = ("Άσκηση " if len(nums) == 1 else "Ασκήσεις ") + " · ".join(str(n) for n in nums)
    run(p, txt, 13, INK, SANS, bold=True)

def subtitle(s, text, y=1261872):
    tf = box(s, L, y, W, 274320, MSO_ANCHOR.MIDDLE)
    rich(tf.paragraphs[0], text, 12, TEXT, SANS)

def footer(s, tag, text):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, L, FOOT_Y, L + 11119104, FOOT_Y)
    ln.line.color.rgb = rgb(LINE); ln.line.width = 9525; ln.line.dash_style = MSO_LINE.DASH
    tf = box(s, L, 6089904, 10241280, 329184, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    run(p, tag + "   ", 9, GOLD, MONO, bold=True, spc=100)
    rich(p, text, 10.5, MUTED, SANS)

KW = set("for in if elif else while break continue pass True False None and or not import from with as try except finally del def return lambda global".split())
BI = set("print input int float str len range sum max min sorted round list tuple abs set dict open enumerate zip map filter help".split())
TOK = re.compile(r'(#.*$)|(f?"[^"]*")|([A-Za-z_]\w*)|(\s+)|(.)')

def code_line(p, line, size):
    if line.startswith("»"):                              # γραμμή εξόδου
        out, _, com = line[1:].partition("   #")
        run(p, out, size, TEXT, MONO)
        if com: run(p, "   #" + com, size, MUTED, MONO)
        return
    for m in TOK.finditer(line):
        com, st, word, sp, ch = m.groups()
        if com: run(p, com, size, MUTED, MONO)
        elif st: run(p, st, size, GOLD, MONO)
        elif word: run(p, word, size, BLUE if (word in KW or word in BI) else INK, MONO)
        else: run(p, sp or ch, size, INK, MONO)
    if not line: run(p, " ", size, INK, MONO)

def code(s, x, y, w, src, bar=BLUE, size=11, tag=None, tag_color=BLUE, h=None):
    """κουτί κώδικα· οι γραμμές που ξεκινούν με » είναι αυτό που εμφανίζεται. Επιστρέφει το κάτω όριο."""
    if tag:
        label(s, x, y, w, tag, tag_color)
        y += 283464
    lines = src.strip("\n").split("\n")
    lh = int(size * 1.2 * 12700)
    hh = h or len(lines) * lh + 2 * 118872
    shape(s, x, y, w, hh, PAPER, LINE)
    shape(s, x, y, 36576, hh, bar)
    tf = box(s, x + 164592, y + 109728, w - 246888, hh - 182880)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(size * 1.2)
        code_line(p, line, size)
    return y + hh

def card(s, x, y, w, title, body, tone="blue", h=None, size=11):
    fill = BLUE_BG if tone == "blue" else GOLD_BG
    bh = est_h(body, w - 329184, size)
    hh = h or (109728 + 274320 + 54864 + bh + 128016)
    shape(s, x, y, w, hh, fill, None, MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box(s, x + 164592, y + 109728, w - 329184, 274320, MSO_ANCHOR.MIDDLE)
    run(tf.paragraphs[0], title, 13, INK, SANS, bold=True)
    tf = box(s, x + 164592, y + 438912, w - 329184, bh)
    for i, part in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        rich(p, part, size, TEXT, SANS, bold_color=INK)
    return y + hh

def cards(s, x, y, w, items, gap=137160):
    for it in items:
        y = card(s, x, y, w, *it) + gap
    return y

def prompt(s, x, y, w, text, tag="ΕΚΦΩΝΗΣΗ", size=11.5):
    bh = est_h(text, w - 219456, size)
    hh = 73152 + 201168 + 36576 + bh + 109728
    shape(s, x, y, w, hh, BLUE_BG, BLUE_BG)
    shape(s, x, y, 41148, hh, BLUE)
    label(s, x + 164592, y + 73152, w - 219456, tag)
    tf = box(s, x + 164592, y + 292608, w - 219456, bh)
    rich(tf.paragraphs[0], text, size, INK, SANS)
    return y + hh

def _borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for i, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = tcPr.makeelement(qn(side), {"w": "9525" if side == "a:lnB" else "0", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        if side == "a:lnB":
            sf = ln.makeelement(qn("a:solidFill"), {}); c = sf.makeelement(qn("a:srgbClr"), {"val": LINE}); sf.append(c); ln.append(sf)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        tcPr.insert(i, ln)

def table(s, x, y, colw, rows, mono=(), size=10.5, row_h=338328, head=True, hl=None):
    """rows[0] = επικεφαλίδες. mono: στήλες σε Consolas. hl: {(γραμμή, στήλη): χρώμα}"""
    gs = s.shapes.add_table(len(rows), len(colw), x, y, sum(colw), row_h * len(rows))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")): tblPr.remove(el)
    for i, cw in enumerate(colw): tbl.columns[i].width = cw
    for r in range(len(rows)):
        tbl.rows[r].height = row_h if (r or not head) else 320040
        for c in range(len(colw)):
            cell = tbl.cell(r, c)
            cell.fill.background()
            _borders(cell)
            cell.margin_left = cell.margin_right = 76200
            cell.margin_top = cell.margin_bottom = 45720
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            txt = str(rows[r][c])
            if r == 0 and head:
                run(p, txt.upper(), 9, MUTED, MONO, bold=True, spc=60)
            else:
                col = (hl or {}).get((r, c), INK)
                rich(p, txt, size, col, MONO if c in mono else SANS)
    return y + sum(tbl.rows[r].height for r in range(len(rows)))

def new(num, kicker, title, lab=None, title_w=None):
    s = prs.slides.add_slide(LAYOUT)
    for el, img in CHROME:
        el = copy.deepcopy(el)
        if img == "bg":                                   # φόντο
            s._element.cSld.insert(0, el)
            continue
        if img is not None:
            el.blipFill.blip.rEmbed = s.part.relate_to(img, RT.IMAGE)
        s.shapes._spTree.append(el)
    kicker_title(s, num, kicker, title, title_w or (8686800 if lab else 9692640))
    if lab: lab_badge(s, lab)
    return s



STEPS = 0
def step(n, name): return f"ΒΗΜΑ {n} ΑΠΟ {STEPS} · {name}"

LVL = {"Εύκολο": "1B7F4B", "Μέτριο": GOLD, "Δύσκολο": "B42318", "Ο στόχος": "B42318", "Για το σπίτι": BLUE}

def lab_overview(s, ex):
    """Πίνακας «οι N ασκήσεις» σε δύο στήλες. ex: [(αρ., τίτλος, επίπεδο, "διαφάνειες"), ...]"""
    half = (len(ex) + 1) // 2
    for k, part in enumerate((ex[:half], ex[half:])):
        rows = [["#", "Άσκηση", "Επίπεδο", "Διαφάνεια"]] + [[str(n), t, lv, d] for n, t, lv, d in part]
        hl = {(i + 1, 2): LVL[part[i][2]] for i in range(len(part))}
        table(s, COL_L if k == 0 else COL_R, TOP, [457200, 2651760, 1234440, 1051560 if k == 0 else 914400], rows,
              mono=(0, 3), size=10.5, row_h=393192 if half <= 9 else 365760, hl=hl)
